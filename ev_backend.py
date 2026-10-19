"""
SAKEC Electric Vehicle Engineering Simulation Platform
Backend Data Persistence, User Management (Admin Only), Assessment Analytics, and Excel Report Engine
Developed by Dr. Rohan Borgalli for Shah & Anchor Kutchhi Engineering College
"""

import os
import json
import time
import hashlib
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
RECORDS_FILE = os.path.join(DATA_DIR, "ev_student_records.json")
USERS_FILE = os.path.join(DATA_DIR, "ev_users.json")

# =========================================================================
# USER AUTHENTICATION & ADMIN ROSTER MANAGEMENT
# =========================================================================

def hash_password(password, salt=None):
    if not salt:
        salt = os.urandom(16).hex()
    hashed = hashlib.sha256((salt + password).encode("utf-8")).hexdigest()
    return hashed, salt

def verify_password(password, salt, hashed):
    test_hash = hashlib.sha256((salt + password).encode("utf-8")).hexdigest()
    return test_hash == hashed

def load_users():
    os.makedirs(DATA_DIR, exist_ok=True)
    if os.path.exists(USERS_FILE):
        try:
            with open(USERS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            print("Error loading users:", e)
    
    # Initialize default admin if file does not exist
    h, s = hash_password("admin@123")
    default_users = {
        "admin@sakec.ac.in": {
            "name": "Dr. Rohan Borgalli",
            "roll": "ADMIN-01",
            "email": "admin@sakec.ac.in",
            "role": "admin",
            "dept": "Computer Engineering",
            "division": "Admin",
            "hash": h,
            "salt": s,
            "created_at": time.time()
        }
    }
    save_users(default_users)
    return default_users

def save_users(users):
    os.makedirs(DATA_DIR, exist_ok=True)
    with open(USERS_FILE, "w", encoding="utf-8") as f:
        json.dump(users, f, indent=2, ensure_ascii=False)

def authenticate_user(email, password):
    users = load_users()
    email_key = email.strip().lower()
    user = users.get(email_key)
    if not user:
        return {"ok": False, "error": "User account not found. Please contact the administrator to enroll you."}
    
    if not verify_password(password, user["salt"], user["hash"]):
        return {"ok": False, "error": "Incorrect password. Please contact the administrator to reset your password."}
    
    return {
        "ok": True,
        "user": {
            "name": user["name"],
            "roll": user.get("roll", "—"),
            "email": user["email"],
            "role": user["role"],
            "dept": user.get("dept", "SAKEC"),
            "division": user.get("division", "—")
        }
    }

def add_user(name, roll, email, role, dept, division="", custom_pass=None):
    users = load_users()
    email_key = email.strip().lower()
    
    # Auto-generate default password: firstname@123 if not provided
    if not custom_pass:
        first_name = name.strip().split()[0].lower() if name.strip() else "sakec"
        # Sanitize first name
        first_name = "".join(c for c in first_name if c.isalnum())
        custom_pass = f"{first_name}@123"
    
    h, s = hash_password(custom_pass)
    users[email_key] = {
        "name": name.strip(),
        "roll": roll.strip().upper(),
        "email": email_key,
        "role": role.strip().lower(), # 'student', 'faculty', 'admin'
        "dept": dept.strip(),
        "division": division.strip(),
        "hash": h,
        "salt": s,
        "default_password": custom_pass,
        "created_at": time.time()
    }
    save_users(users)
    return users[email_key], custom_pass

def reset_user_password(email, new_password=None):
    users = load_users()
    email_key = email.strip().lower()
    if email_key not in users:
        return None, "User not found."
    
    u = users[email_key]
    if not new_password:
        first_name = u["name"].strip().split()[0].lower() if u["name"].strip() else "sakec"
        first_name = "".join(c for c in first_name if c.isalnum())
        new_password = f"{first_name}@123"
        
    h, s = hash_password(new_password)
    u["hash"] = h
    u["salt"] = s
    u["default_password"] = new_password
    save_users(users)
    return new_password, None

def delete_user(email):
    users = load_users()
    email_key = email.strip().lower()
    if email_key in users and email_key != "admin@sakec.ac.in":
        del users[email_key]
        save_users(users)
        return True
    return False

def bulk_upload_users_csv(csv_text, default_role="student"):
    lines = csv_text.strip().split("\n")
    created = 0
    updated = 0
    details = []

    for line in lines:
        parts = [p.strip() for p in line.split(",") if p.strip()]
        if not parts or parts[0].lower() in ["email", "email id", "roll", "roll no"]:
            continue
        
        # Expected formats:
        # Roll, Name, Email, Dept, Division
        # OR Email, Name, Dept, Roll
        if len(parts) >= 3:
            if "@" in parts[0]:
                email = parts[0]
                name = parts[1]
                dept = parts[2]
                roll = parts[3] if len(parts) > 3 else "—"
                div = parts[4] if len(parts) > 4 else "—"
            else:
                roll = parts[0]
                name = parts[1]
                email = parts[2]
                dept = parts[3] if len(parts) > 3 else "SAKEC"
                div = parts[4] if len(parts) > 4 else "—"

            u, p = add_user(name, roll, email, default_role, dept, div)
            created += 1
            details.append({"email": email, "name": name, "roll": roll, "pass": p})
            
    return created, details

# =========================================================================
# SIMULATION PROGRESS & ASSESSMENT ENGINE
# =========================================================================

def load_records():
    if os.path.exists(RECORDS_FILE):
        try:
            with open(RECORDS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            print("Error loading records:", e)
    return {
        "students": {},
        "metadata": {
            "college": "Shah & Anchor Kutchhi Engineering College, Chembur, Mumbai",
            "faculty": "Dr. Rohan Borgalli"
        }
    }

def save_records(records):
    os.makedirs(DATA_DIR, exist_ok=True)
    with open(RECORDS_FILE, "w", encoding="utf-8") as f:
        json.dump(records, f, indent=2, ensure_ascii=False)

def complete_simulation(email, name, roll, dept, division, module_id, reflection, score, total_q=5, params=None):
    records = load_records()
    email_key = email.strip().lower()
    
    if email_key not in records["students"]:
        records["students"][email_key] = {
            "name": name.strip(),
            "roll": roll.strip().upper(),
            "email": email_key,
            "dept": dept.strip(),
            "division": division.strip(),
            "registered_at": time.time(),
            "last_active": time.time(),
            "completed_modules": {}
        }

    st_data = records["students"][email_key]
    prev_module_data = st_data["completed_modules"].get(module_id, {})
    prev_attempts = prev_module_data.get("attempts", 0)
    prev_best_score = prev_module_data.get("best_score", 0)

    # STRICT PASS THRESHOLD: Must get >= 3 of 5 correct (>= 60%)
    passed = (score >= 3)
    best_score = max(prev_best_score, score)

    st_data["completed_modules"][module_id] = {
        "completed": passed,
        "status": "COMPLETED" if passed else "RETRY_REQUIRED",
        "last_attempt_at": time.time(),
        "reflection": reflection.strip(),
        "quiz_score": int(score),
        "quiz_total": int(total_q),
        "quiz_percentage": round((score / total_q) * 100, 1),
        "best_score": int(best_score),
        "attempts": prev_attempts + 1,
        "parameters": params or {}
    }
    st_data["last_active"] = time.time()
    save_records(records)
    return passed

def get_student_profile(email):
    records = load_records()
    return records["students"].get(email.strip().lower(), None)

def get_all_students_progress():
    records = load_records()
    users = load_users()
    
    students_list = []
    
    # Show all enrolled students from users database
    for email, u in users.items():
        if u.get("role") != "student":
            continue
        
        st_record = records["students"].get(email, {})
        comp = st_record.get("completed_modules", {})
        
        # Completed means module passed (>= 3/5)
        completed_modules_count = len([m for m in comp.values() if m.get("completed") is True])
        scores = [m.get("quiz_percentage", 0) for m in comp.values() if m.get("completed") is True]
        avg_score = round(sum(scores) / len(scores), 1) if scores else 0.0

        students_list.append({
            "roll": u.get("roll", "—"),
            "name": u.get("name", "—"),
            "email": email,
            "dept": u.get("dept", "—"),
            "division": u.get("division", "—"),
            "completed_count": completed_modules_count,
            "total_modules": 7,
            "completion_pct": round((completed_modules_count / 7) * 100, 1),
            "average_score": avg_score,
            "last_active": st_record.get("last_active", u.get("created_at", 0)),
            "modules": comp
        })

    students_list.sort(key=lambda x: x["roll"])
    return students_list

def export_master_excel(output_path):
    students = get_all_students_progress()
    wb = openpyxl.Workbook()

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    navy_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    green_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
    orange_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    title_font = Font(name="Calibri", size=15, bold=True, color="1B365D")
    sub_font = Font(name="Calibri", size=10, italic=True, color="555555")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    ws1 = wb.active
    ws1.title = "EV_Progress_Register"
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = "SHAH & ANCHOR KUTCHHI ENGINEERING COLLEGE, CHEMBUR, MUMBAI"
    ws1["A1"].font = title_font
    ws1["A2"] = "Electric Vehicle Engineering Simulation Platform — Class Progress Register (NBA Criterion 2.2.1)"
    ws1["A2"].font = bold_font
    ws1["A3"] = "Developed by Dr. Rohan Borgalli · Pass Rule: Minimum 3 of 5 Correct Answers Required"
    ws1["A3"].font = sub_font

    headers = [
        "Roll No", "Student Name", "Department", "Division",
        "Mod 1 (Dynamics)", "Mod 2 (Motors)", "Mod 3 (Battery)",
        "Mod 4 (Regen)", "Mod 5 (Charging)", "Mod 6 (Decarbon)",
        "Mod 7 (TCO)", "Completed / 7", "Avg Score (%)", "Overall Status"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_num = 6
    module_keys = [f"module_{i}" for i in range(1, 8)]

    for st in students:
        ws1.cell(row=row_num, column=1, value=st["roll"]).font = bold_font
        ws1.cell(row=row_num, column=2, value=st["name"]).font = normal_font
        ws1.cell(row=row_num, column=3, value=st["dept"]).font = normal_font
        ws1.cell(row=row_num, column=4, value=st["division"]).font = normal_font

        for m_idx, m_key in enumerate(module_keys, 5):
            comp_data = st["modules"].get(m_key, {})
            is_comp = comp_data.get("completed", False)
            score = comp_data.get("quiz_score", 0)
            attempts = comp_data.get("attempts", 0)
            cell = ws1.cell(row=row_num, column=m_idx)
            if is_comp:
                cell.value = f"✓ Pass ({score}/5)"
                cell.fill = green_fill
                cell.font = Font(name="Calibri", size=9, color="0F5132", bold=True)
            elif attempts > 0:
                cell.value = f"✗ Retry ({score}/5)"
                cell.fill = orange_fill
                cell.font = Font(name="Calibri", size=9, color="856404")
            else:
                cell.value = "—"
                cell.font = Font(name="Calibri", size=9, color="888888")
            cell.alignment = Alignment(horizontal="center")

        ws1.cell(row=row_num, column=12, value=f"{st['completed_count']} / 7").alignment = Alignment(horizontal="center")
        ws1.cell(row=row_num, column=13, value=f"{st['average_score']}%").alignment = Alignment(horizontal="center")
        
        status_cell = ws1.cell(row=row_num, column=14)
        if st["completed_count"] == 7:
            status_cell.value = "ALL COMPLETED"
            status_cell.fill = green_fill
            status_cell.font = Font(name="Calibri", size=9.5, bold=True, color="0F5132")
        elif st["completed_count"] >= 4:
            status_cell.value = "IN PROGRESS"
            status_cell.fill = orange_fill
        else:
            status_cell.value = "PENDING"
        status_cell.alignment = Alignment(horizontal="center")

        for c in range(1, 15):
            ws1.cell(row=row_num, column=c).border = thin_border
        row_num += 1

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 13)

    # SHEET 2: Student Reflections Dossier
    ws2 = wb.create_sheet(title="Student_Reflections")
    ws2.views.sheetView[0].showGridLines = True
    ws2["A1"] = "STUDENT TECHNICAL REFLECTION DOSSIER"
    ws2["A1"].font = title_font
    ws2["A2"] = "Qualitative Engineering Observations Submitted per Module"
    ws2["A2"].font = sub_font

    ref_headers = ["Roll No", "Student Name", "Module", "Score (/5)", "Attempts", "Status", "Submitted Reflection"]
    for col_idx, h in enumerate(ref_headers, 1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ref_row = 5
    for st in students:
        for m_key, m_data in st["modules"].items():
            ws2.cell(row=ref_row, column=1, value=st["roll"]).font = bold_font
            ws2.cell(row=ref_row, column=2, value=st["name"]).font = normal_font
            ws2.cell(row=ref_row, column=3, value=m_key.replace("module_", "Module ")).font = normal_font
            ws2.cell(row=ref_row, column=4, value=f"{m_data.get('quiz_score', 0)}/5").alignment = Alignment(horizontal="center")
            ws2.cell(row=ref_row, column=5, value=m_data.get("attempts", 1)).alignment = Alignment(horizontal="center")
            
            stat_c = ws2.cell(row=ref_row, column=6, value=m_data.get("status", "—"))
            stat_c.alignment = Alignment(horizontal="center")
            if m_data.get("completed"):
                stat_c.fill = green_fill
            else:
                stat_c.fill = orange_fill
                
            ref_c = ws2.cell(row=ref_row, column=7, value=m_data.get("reflection", "—"))
            ref_c.alignment = Alignment(wrap_text=True)
            for c in range(1, 8):
                ws2.cell(row=ref_row, column=c).border = thin_border
            ref_row += 1

    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 22
    ws2.column_dimensions["C"].width = 16
    ws2.column_dimensions["D"].width = 12
    ws2.column_dimensions["E"].width = 12
    ws2.column_dimensions["F"].width = 16
    ws2.column_dimensions["G"].width = 80

    wb.save(output_path)
    return output_path
